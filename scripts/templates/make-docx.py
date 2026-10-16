"""Word templates built with python-docx: classic single column, contact in the page header,
a two-column table layout (the most common Word template), and a section with two newspaper
columns. build-templates.ps1 then adds a floating text box via Word itself and exports every
.docx to PDF with Word, so the PDFs are genuine Word output."""
import json
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

P = json.load(open("fixtures/templates/person.json", encoding="utf8"))
OUT = "fixtures/templates/docx"
os.makedirs(OUT, exist_ok=True)


def base(font="Calibri", size=11):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = font
    st.font.size = Pt(size)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(0.75)
        s.top_margin = s.bottom_margin = Inches(0.7)
    return d


def heading(d_or_cell, text):
    p = d_or_cell.add_paragraph()
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(12)
    return p


def roles(container):
    for r in P["roles"]:
        p = container.add_paragraph()
        p.add_run(f"{r['title']} | {r['company']}").bold = True
        p.add_run(f"\t{r['dates']}")
        for b in r["bullets"]:
            container.add_paragraph(b, style="List Bullet")


def body(d, contact=True):
    t = d.add_paragraph()
    run = t.add_run(P["name"])
    run.bold = True
    run.font.size = Pt(20)
    if contact:
        d.add_paragraph(f"{P['location']} | {P['email']} | {P['phone']} | {P['linkedin']}")
    heading(d, "Professional Summary")
    d.add_paragraph(P["summary"])
    heading(d, "Work Experience")
    roles(d)
    heading(d, "Education")
    d.add_paragraph(P["education"])
    heading(d, "Skills")
    d.add_paragraph(", ".join(P["skills"]))


# 6. Classic single column (control: should come out clean).
d = base()
body(d)
d.save(f"{OUT}/word-classic.docx")

# 7. Contact details in the Word page header.
d = base("Arial", 10.5)
d.sections[0].header.paragraphs[0].text = f"{P['email']}  |  {P['phone']}  |  {P['linkedin']}"
d.sections[0].header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
body(d, contact=False)
d.save(f"{OUT}/word-header-contact.docx")

# 8. Two-column table layout: left cell = contact/skills/education, right cell = experience.
d = base("Calibri", 10.5)
name = d.add_paragraph()
r = name.add_run(P["name"])
r.bold = True
r.font.size = Pt(22)
tbl = d.add_table(rows=1, cols=2)
tbl.autofit = False
left, right = tbl.rows[0].cells
left.width, right.width = Inches(2.3), Inches(4.7)
left.paragraphs[0].text = "CONTACT"
for line in [P["email"], P["phone"], P["location"], P["linkedin"]]:
    left.add_paragraph(line)
heading(left, "Skills")
for s in P["skills"]:
    left.add_paragraph(s)
heading(left, "Education")
left.add_paragraph(P["education"])
right.paragraphs[0].text = "PROFILE"
right.add_paragraph(P["summary"])
heading(right, "Experience")
roles(right)
d.save(f"{OUT}/word-table-layout.docx")

# 9. Newspaper columns (Layout > Columns > Two) for the whole document.
d = base("Cambria", 10.5)
body(d)
cols = OxmlElement("w:cols")
cols.set(qn("w:num"), "2")
cols.set(qn("w:space"), "360")
d.sections[0]._sectPr.append(cols)
d.save(f"{OUT}/word-two-columns.docx")

# 10. Classic again; build-templates.ps1 adds a floating text box with the contact block.
d = base("Georgia", 10.5)
body(d, contact=False)
d.save(f"{OUT}/word-textbox.docx")
print("docx templates written")

# 11. A realistic "needs work" Word resume: creative headings, weak bullets, mixed dates,
# bold role lines and real Word bullets. Used to test in-place tailoring.
d = base("Calibri", 11)
t = d.add_paragraph()
r = t.add_run("Priya Shah")
r.bold = True
r.font.size = Pt(20)
d.add_paragraph("priya.shah@example.com | +1 (415) 555-0199 | San Francisco, CA")
heading(d, "About Me")
d.add_paragraph("Product manager with 7 years of experience shipping B2B SaaS products and growth experiments.")
heading(d, "Where I've Been")
p = d.add_paragraph()
p.add_run("Senior Product Manager | Northwind Software").bold = True
p.add_run("\t01/2021 - Present")
for b in [
    "Owned the roadmap for the billing product line, growing net revenue retention from 104% to 117%.",
    "Responsible for writing PRDs and user stories for the engineering team.",
    "Worked on A/B tests to improve conversion.",
]:
    d.add_paragraph(b, style="List Bullet")
p = d.add_paragraph()
p.add_run("Product Manager | Lumen Retail Tech").bold = True
p.add_run("\tSummer 2018 - Dec 2020")
for b in ["Managed the mobile checkout experience for 2M monthly users.", "Helped the sales team with enterprise demos."]:
    d.add_paragraph(b, style="List Bullet")
heading(d, "Education")
d.add_paragraph("MBA, Haas School of Business, UC Berkeley")
heading(d, "What I Know")
d.add_paragraph("Roadmapping, Jira, Figma, SQL, A/B Testing, Agile")
d.save(f"{OUT}/word-needs-work.docx")
print("word-needs-work written")
